import openpyxl
import psycopg2
from psycopg2.extras import execute_values
import os
import sys

# Forçar saída UTF-8 no Windows
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

def main():
    excel_path = r'D:\IA MALELO\banco_dados\BD MINIPCP.xlsx'
    connection_string = 'postgresql://neondb_owner:npg_EcWR7JP8tNGK@ep-cold-fire-acm5ge04.sa-east-1.aws.neon.tech/neondb?sslmode=require'

    print(f"Lendo arquivo Excel: {excel_path}...")
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    sheet = wb.active

    rows_to_insert = []
    for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not any(row):
            continue
        
        tipo = str(row[0]).strip() if row[0] is not None else 'Indefinido'
        codigo = str(row[1]).strip() if row[1] is not None else ''
        descricao = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ''

        if not tipo and not codigo and not descricao:
            continue

        # Normalizar categoria
        categoria = tipo
        if 'aço' in tipo.lower() or 'aco' in tipo.lower():
            categoria = 'Aço'
        elif 'alumin' in tipo.lower():
            categoria = 'Alumínio'
        elif 'lat' in tipo.lower():
            categoria = 'Latão'
        elif 'bronze' in tipo.lower():
            categoria = 'Bronze'
        elif 'plastic' in tipo.lower() or 'polimer' in tipo.lower():
            categoria = 'Plásticos'
        elif 'produto' in tipo.lower() and 'acabado' in tipo.lower():
            categoria = 'Produto Acabado'
        elif 'forneced' in tipo.lower():
            categoria = 'Fornecedor'

        rows_to_insert.append((tipo, codigo, descricao, categoria))

    print(f"Total de {len(rows_to_insert)} itens carregados do Excel!")

    print("Conectando ao Neon PostgreSQL (orcamentista lasec - sa-east-1)...")
    conn = psycopg2.connect(connection_string)
    cur = conn.cursor()

    print("Criando tabela 'minipcp_itens' no Neon...")
    cur.execute("""
        CREATE TABLE IF NOT EXISTS minipcp_itens (
            id SERIAL PRIMARY KEY,
            tipo VARCHAR(100) NOT NULL,
            codigo VARCHAR(100),
            descricao TEXT NOT NULL,
            categoria VARCHAR(100),
            created_at TIMESTAMP WITH TIME ZONE DEFAULT CURRENT_TIMESTAMP
        );

        CREATE INDEX IF NOT EXISTS idx_minipcp_tipo ON minipcp_itens(tipo);
        CREATE INDEX IF NOT EXISTS idx_minipcp_codigo ON minipcp_itens(codigo);
        CREATE INDEX IF NOT EXISTS idx_minipcp_categoria ON minipcp_itens(categoria);
    """)

    # Limpar tabela antes de inserir para evitar duplicação
    cur.execute("TRUNCATE TABLE minipcp_itens RESTART IDENTITY;")

    print(f"Inserindo {len(rows_to_insert)} registros em lote no Neon...")
    insert_query = """
        INSERT INTO minipcp_itens (tipo, codigo, descricao, categoria)
        VALUES %s
    """
    execute_values(cur, insert_query, rows_to_insert, page_size=500)
    conn.commit()

    # Verificação
    cur.execute("SELECT count(*) FROM minipcp_itens;")
    total_count = cur.fetchone()[0]

    cur.execute("""
        SELECT categoria, count(*) 
        FROM minipcp_itens 
        GROUP BY categoria 
        ORDER BY count(*) DESC 
        LIMIT 10;
    """)
    top_categories = cur.fetchall()

    cur.execute("""
        SELECT table_name 
        FROM information_schema.tables 
        WHERE table_schema = 'public' 
        ORDER BY table_name;
    """)
    tables = cur.fetchall()

    print("\n" + "="*50)
    print(f"SUCESSO! Tabela 'minipcp_itens' criada e populada com {total_count} itens!")
    print("="*50)
    print("Distribuicao por Categorias:")
    for cat, cnt in top_categories:
        print(f"  • {cat}: {cnt} itens")
    
    print("\nTodas as Tabelas Ativas no seu Neon:")
    for t in tables:
        print(f"  • {t[0]}")

    cur.close()
    conn.close()

if __name__ == '__main__':
    main()